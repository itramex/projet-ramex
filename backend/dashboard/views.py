from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from django.db.models import Count, Sum, F, Q, Case, When, Value, CharField
from django.db.models.functions import TruncMonth
from django.http import HttpResponse
from datetime import datetime, timedelta
from .models import VillageReference

from producteurs.models import Producteur, AGR, Dotation
from parcelles.models import Parcelle
from tracabilite.models import BonCollecte, LotTraitement, Colis, CommandeExport, TracabiliteChain


def _dashboard_base_queryset(request):
    """Queryset Producteur actif filtré par village/commune/fokontany (query params GET).

    Réutilisé par les endpoints du dashboard (hygiene, environnement, enfants…)
    pour que les filtres du frontend s'appliquent à tous les onglets.
    """
    qs = Producteur.objects.filter(actif=True)
    villages = [v for v in request.query_params.getlist('village') if v]
    if villages:
        qs = qs.filter(village__in=villages)
    communes = [c for c in request.query_params.getlist('commune') if c]
    if communes:
        qs = qs.filter(commune__in=communes)
    fokontanys = [f for f in request.query_params.getlist('fokontany') if f]
    if fokontanys:
        qs = qs.filter(fokontany__in=fokontanys)

    # #29 — Confidentialité par agence : les non-responsables ne voient que
    # les producteurs de leur agence (via la coopérative).
    from users.permissions import scope_par_agence
    qs, _ = scope_par_agence(request.user, qs)
    return qs


def _children_schooling_stats(queryset, ref_year):
    """Statistiques de scolarisation robustes (cascade de sources).

    Contexte terrain (mesuré sur la base réelle) :
      - les compteurs d'enfants déclarés sont fiables (nb_enfants_garcons/filles…)
      - `nb_enfants_scolarises` n'est PAS rempli par l'import historique (souvent 0)
      - `niveau_etude_enfant_N` est vide
      - `annee_naissance_enfant_N` n'est rempli qu'à ~32 % et contient parfois des
        années d'adultes (ex. 1980) : le calcul par slots ne retrouve que ~16 enfants
        sur 466 déclarés.

    On applique donc une cascade, du plus fiable au moins fiable :
      1. `nb_enfants_scolarises` importé (> 0) → source directe
      2. sinon : compteurs d'enfants déclarés moins ceux marqués non scolarisés
      3. sinon : calcul par slots (année de naissance 3-18 ans + continue_ecole)

    Retourne un dict (et non plus un tuple) pour exposer la méthode utilisée.
    """
    child_fields = [f'annee_naissance_enfant_{i}' for i in range(1, 11)]
    school_fields = [f'continue_ecole_enfant_{i}' for i in range(1, 11)]

    slots_age = 0
    slots_schooled = 0
    declared_children = 0
    declared_schooled = 0
    declared_not_schooled = 0

    rows = queryset.values_list(
        *child_fields,
        *school_fields,
        'nb_enfants_garcons',
        'nb_enfants_filles',
        'nb_autres_garcons',
        'nb_autres_filles',
        'nb_enfants_scolarises',
        'nb_enfants_non_scolarises',
    )
    for row in rows:
        for i in range(10):
            annee_naissance = row[i]
            if not annee_naissance:
                continue
            age = ref_year - annee_naissance
            if 3 <= age <= 18:
                slots_age += 1
                if row[10 + i]:
                    slots_schooled += 1

        # Accumulation sur TOUTES les lignes (et non la dernière uniquement)
        declared_children += (row[20] or 0) + (row[21] or 0) + (row[22] or 0) + (row[23] or 0)
        declared_schooled += row[24] or 0
        declared_not_schooled += row[25] or 0

    # --- Cascade ---
    if declared_schooled > 0:
        methode = 'declare_agrege'
        scolarises = declared_schooled
        non_scolarises = declared_not_schooled
    elif declared_children > 0:
        # Aucun champ agrégé fiable : les enfants NON signalés comme non scolarisés
        # sont considérés scolarisés (les signalements non scolarisés sont explicites).
        methode = 'deduit_compteurs'
        scolarises = max(0, declared_children - declared_not_schooled)
        non_scolarises = declared_not_schooled
    else:
        methode = 'slots'
        scolarises = slots_schooled
        non_scolarises = max(0, slots_age - slots_schooled)

    en_age_scolaire = declared_children if declared_children > 0 else slots_age

    return {
        'total_enfants': declared_children,
        'enfants_en_age_scolaire': en_age_scolaire,
        'enfants_scolarises': scolarises,
        'enfants_non_scolarises': non_scolarises,
        'taux_scolarisation': round((scolarises / en_age_scolaire * 100) if en_age_scolaire > 0 else 0, 2),
        'methode': methode,
        'slots_age': slots_age,
        'slots_scolarises': slots_schooled,
    }


def _children_school_age_and_schooled(queryset, ref_year):
    """Compatibilité : ancien tuple (age, scolarises) dérivé du helper robuste."""
    stats = _children_schooling_stats(queryset, ref_year)
    return stats['enfants_en_age_scolaire'], stats['enfants_scolarises']



@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_global(request):
    """
    Vue d'ensemble globale du dashboard
    GET /api/dashboard/
    
    Paramètres de filtrage:
    - village: filtre par village(s) - peut être multiple
    - commune: filtre par commune(s) - peut être multiple
    - year: année pour les données historiques (optionnel)
    """
    # Récupérer les paramètres de filtrage (support multi-sélection)
    villages = request.query_params.getlist('village', None)
    communes = request.query_params.getlist('commune', None)
    fokontanys = request.query_params.getlist('fokontany', None)
    year = request.query_params.get('year', None)
    try:
        selected_year = int(year) if year else None
    except ValueError:
        selected_year = None
    
    # Construire le queryset de base avec les filtres
    base_queryset = Producteur.objects.filter(actif=True).prefetch_related('dotations')
    base_queryset_all = Producteur.objects.all()
    
    if villages and any(villages):
        # Filtrer les valeurs vides
        villages = [v for v in villages if v]
        if villages:
            base_queryset = base_queryset.filter(village__in=villages)
            base_queryset_all = base_queryset_all.filter(village__in=villages)
    
    if communes and any(communes):
        # Filtrer les valeurs vides
        communes = [c for c in communes if c]
        if communes:
            base_queryset = base_queryset.filter(commune__in=communes)
            base_queryset_all = base_queryset_all.filter(commune__in=communes)

    if fokontanys and any(fokontanys):
        fokontanys = [f for f in fokontanys if f]
        if fokontanys:
            base_queryset = base_queryset.filter(fokontany__in=fokontanys)
            base_queryset_all = base_queryset_all.filter(fokontany__in=fokontanys)

    # #29 — Confidentialité par agence : les non-responsables (animateur,
    # agent de collecte) ne voient que les producteurs de leur agence
    # (via la coopérative). Admin/superviseur → tout.
    from users.permissions import scope_par_agence
    base_queryset, _ = scope_par_agence(request.user, base_queryset)
    base_queryset_all, _ = scope_par_agence(request.user, base_queryset_all)

    # Statistiques globales (sans filtres pour le total)
    # TOTAL = TOUS LES PRODUCTEURS (actifs + inactifs) — #29 : restreint aux
    # producteurs visibles par l'utilisateur (scopés à son agence).
    _scoped_all = scope_par_agence(request.user, Producteur.objects.all())[0]
    total = _scoped_all.count()  # Tous les producteurs visibles
    actifs_sans_filtres = scope_par_agence(
        request.user, Producteur.objects.filter(actif=True))[0].count()
    inactifs = scope_par_agence(
        request.user, Producteur.objects.filter(actif=False))[0].count()
    
    # Actifs avec filtres appliqués (pour statistiques détaillées)
    actifs = base_queryset.count()
    total_filtres = base_queryset_all.count()
    inactifs_filtres = base_queryset_all.filter(actif=False).count()
    
    # Statistiques par genre
    par_genre = list(
        base_queryset
        .values('sexe')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    
    # Répartition par âge et sexe — agrégation DB au lieu de boucle Python
    today = datetime.now().date()
    birth_cutoff_25 = today.replace(year=today.year - 25)
    birth_cutoff_35 = today.replace(year=today.year - 35)
    birth_cutoff_45 = today.replace(year=today.year - 45)

    age_data = (
        base_queryset
        .exclude(date_naissance__isnull=True)
        .annotate(
            age_group=Case(
                When(date_naissance__gt=birth_cutoff_25, then=Value('<25')),
                When(date_naissance__gt=birth_cutoff_35, then=Value('25-34')),
                When(date_naissance__gt=birth_cutoff_45, then=Value('35-44')),
                default=Value('45+'),
                output_field=CharField(),
            )
        )
        .values('age_group', 'sexe')
        .annotate(count=Count('id'))
    )

    age_distribution = {k: {'hommes': 0, 'femmes': 0} for k in ['<25', '25-34', '35-44', '45+']}
    for row in age_data:
        group = row['age_group']
        if row['sexe'] == 'M':
            age_distribution[group]['hommes'] += row['count']
        elif row['sexe'] == 'F':
            age_distribution[group]['femmes'] += row['count']

    # Calcul des pourcentages par rapport au total actifs
    total_actifs = actifs or 0
    age_distribution_list = []
    for label in ['<25', '25-34', '35-44', '45+']:
        hommes = age_distribution[label]['hommes']
        femmes = age_distribution[label]['femmes']
        age_distribution_list.append({
            'groupe': label,
            'hommes': hommes,
            'femmes': femmes,
            'hommes_pct': round((hommes / total_actifs * 100), 2) if total_actifs > 0 else 0,
            'femmes_pct': round((femmes / total_actifs * 100), 2) if total_actifs > 0 else 0,
            'hommes_abs': hommes,  # Nombre absolu pour l'axe Y
            'femmes_abs': femmes,  # Nombre absolu pour l'axe Y
        })
    
    # Top 10 communes
    top_communes = list(
        base_queryset
        .values('commune')
        .annotate(count=Count('id'))
        .order_by('-count')[:10]
    )
    
    # Par coopérative
    par_cooperative = list(
        base_queryset.filter(cooperative__isnull=False)
        .values('cooperative__nom')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    
    # Niveau d'éducation
    par_education = list(
        base_queryset
        .values('niveau_education')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    
    # Statut matrimonial
    par_statut = list(
        base_queryset
        .values('statut_matrimonial')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    
    # Leadership
    femmes_leaders = base_queryset.filter(femme_leader=True).count()
    paysans_relais = base_queryset.filter(paysan_relais=True).count()
    
    # Responsabilités coopérative (non "aucune")
    responsables = base_queryset.exclude(responsabilite_cooperative='aucune').count()
    
    # Statistiques enfants
    # Le total d'enfants est TOUJOURS calculé depuis les données actuelles du modèle
    agregats_enfants = base_queryset.aggregate(
        total_enfants=Sum(
            F('nb_enfants_garcons') + F('nb_enfants_filles') + 
            F('nb_autres_garcons') + F('nb_autres_filles')
        ),
    )
    total_enfants = agregats_enfants['total_enfants'] or 0
    
    # Pour la scolarisation, utiliser l'historique seulement pour les années passées
    # quand des données historiques existent réellement
    use_current_data = True
    annee_actuelle = datetime.now().year
    reference_year = selected_year or annee_actuelle
    
    if year:
        try:
            from history.models import SocialIndicatorHistory
            year_int = int(year)
            
            # Utiliser l'historique seulement pour les années passées
            if year_int < annee_actuelle:
                # Construire le queryset d'historique avec les mêmes filtres
                social_history_qs = SocialIndicatorHistory.objects.filter(annee=year_int)
                
                if villages and any(villages):
                    villages_clean = [v for v in villages if v]
                    if villages_clean:
                        social_history_qs = social_history_qs.filter(producteur__village__in=villages_clean)
                
                if communes and any(communes):
                    communes_clean = [c for c in communes if c]
                    if communes_clean:
                        social_history_qs = social_history_qs.filter(producteur__commune__in=communes_clean)
                
                if fokontanys and any(fokontanys):
                    fokontanys_clean = [f for f in fokontanys if f]
                    if fokontanys_clean:
                        social_history_qs = social_history_qs.filter(producteur__fokontany__in=fokontanys_clean)
                
                # Récupérer les indicateurs de scolarisation pour l'année
                scolarisation_records = social_history_qs.filter(type_indicateur='scolarisation')
                
                # Utiliser l'historique seulement si des enregistrements existent
                if scolarisation_records.exists():
                    use_current_data = False
                    enfants_scolarises = 0
                    for record in scolarisation_records:
                        if record.valeur_numerique is not None:
                            enfants_scolarises += int(record.valeur_numerique)
                    enfants_non_scolarises = max(0, total_enfants - enfants_scolarises)
                    enfants_en_age_scolaire = total_enfants
                    
        except (ValueError, ImportError) as e:
            print(f"Erreur lors de la récupération des données historiques: {e}")
            use_current_data = True
    
    # Valeurs par défaut (cas « données historiques » sans reconstruction des slots)
    scolarisation_methode = 'historique'
    scolarisation_slots = {'age': 0, 'scolarises': 0}

    # Utiliser les données actuelles si pas d'historique disponible
    if use_current_data:
        schooling = _children_schooling_stats(base_queryset, reference_year)
        enfants_en_age_scolaire = schooling['enfants_en_age_scolaire']
        enfants_scolarises = schooling['enfants_scolarises']
        enfants_non_scolarises = schooling['enfants_non_scolarises']
        scolarisation_methode = schooling['methode']
        scolarisation_slots = {
            'age': schooling['slots_age'],
            'scolarises': schooling['slots_scolarises'],
        }

    # Impact des kits scolaires et statistiques par tranche d'âge
    # Seulement pour les données actuelles (pas pour l'historique)
    if use_current_data:
        # ⚠️ NE PAS réinitialiser enfants_en_age_scolaire ici : la valeur vient du
        # helper robuste (cascade). Avant, elle était remise à 0 puis recalculée
        # uniquement depuis les slots annee_naissance_enfant_N (~32 % remplis et
        # parfois aberrants), d'où l'affichage « 0 enfant scolarisé » (#34).
        # On garde donc un compteur SÉPARÉ pour les slots.
        slots_en_age_scolaire = 0

        # children_school_age: nombre d'enfants 3-18 ans avec année de naissance connue
        # schooled_estimated: estimation des enfants scolarisés parmi les 3-18 ans (exclut sans année de naissance)
        impact_stats = {
            'with_kit': {'children_school_age': 0, 'schooled_estimated': 0},
            'without_kit': {'children_school_age': 0, 'schooled_estimated': 0}
        }
        
        # Statistiques par tranche d'âge pour la pyramide de scolarisation
        age_groups = {
            '3-6': {'total': 0, 'scolarises': 0, 'producteurs': []},
            '7-10': {'total': 0, 'scolarises': 0, 'producteurs': []},
            '11-14': {'total': 0, 'scolarises': 0, 'producteurs': []},
            '15-18': {'total': 0, 'scolarises': 0, 'producteurs': []},
        }
        
        # Première passe: compter tous les enfants par tranche d'âge
        for producteur in base_queryset:
            enfants_producteur_par_age = {'3-6': 0, '7-10': 0, '11-14': 0, '15-18': 0}
            
            # Vérifier si le producteur a reçu un kit scolaire
            has_kit = any(d.type_dotation == 'kit_scolaire' for d in producteur.dotations.all())

            # Préparer compteurs par producteur
            nb_scolarises_producteur = producteur.nb_enfants_scolarises or 0
            known_total_any = 0
            known_3_18 = 0

            for i in range(1, 11):  # annee_naissance_enfant_1 à annee_naissance_enfant_10
                annee_naissance = getattr(producteur, f'annee_naissance_enfant_{i}', None)
                if annee_naissance:
                    known_total_any += 1
                    age = reference_year - annee_naissance
                    if 3 <= age <= 18:
                        slots_en_age_scolaire += 1
                        known_3_18 += 1

                        if has_kit:
                            impact_stats['with_kit']['children_school_age'] += 1
                        else:
                            impact_stats['without_kit']['children_school_age'] += 1
                        
                        # Déterminer la tranche d'âge
                        if 3 <= age <= 6:
                            age_groups['3-6']['total'] += 1
                            enfants_producteur_par_age['3-6'] += 1
                        elif 7 <= age <= 10:
                            age_groups['7-10']['total'] += 1
                            enfants_producteur_par_age['7-10'] += 1
                        elif 11 <= age <= 14:
                            age_groups['11-14']['total'] += 1
                            enfants_producteur_par_age['11-14'] += 1
                        elif 15 <= age <= 18:
                            age_groups['15-18']['total'] += 1
                            enfants_producteur_par_age['15-18'] += 1
        
            # Estimer le nombre d'enfants scolarisés 3-18 ans pour l'impact (exclut sans année de naissance)
            if known_total_any > 0 and known_3_18 > 0 and nb_scolarises_producteur > 0:
                est_schooled_3_18 = nb_scolarises_producteur * (known_3_18 / known_total_any)
                if has_kit:
                    impact_stats['with_kit']['schooled_estimated'] += est_schooled_3_18
                else:
                    impact_stats['without_kit']['schooled_estimated'] += est_schooled_3_18

            # Stocker les infos du producteur avec ses enfants par âge
            for age_range in age_groups.keys():
                if enfants_producteur_par_age[age_range] > 0:
                    age_groups[age_range]['producteurs'].append({
                        'producteur': producteur,
                        'nb_enfants_dans_tranche': enfants_producteur_par_age[age_range]
                    })
    
        # Deuxième passe: répartir les enfants scolarisés proportionnellement
        for age_range, data in age_groups.items():
            total_enfants_tranche = data['total']
            
            if total_enfants_tranche > 0:
                # Pour chaque producteur ayant des enfants dans cette tranche
                for prod_info in data['producteurs']:
                    producteur = prod_info['producteur']
                    nb_enfants_prod_tranche = prod_info['nb_enfants_dans_tranche']
                    
                    # Calculer combien d'enfants de ce producteur sont dans cette tranche d'âge
                    # Et répartir proportionnellement les enfants scolarisés
                    total_enfants_producteur = producteur.nb_enfants_garcons + producteur.nb_enfants_filles
                    nb_scolarises_producteur = producteur.nb_enfants_scolarises or 0
                    
                    if total_enfants_producteur > 0:
                        # Proportion d'enfants de ce producteur qui sont dans cette tranche
                        proportion_dans_tranche = nb_enfants_prod_tranche / total_enfants_producteur
                        # Nombre estimé de scolarisés dans cette tranche
                        scolarises_estimes = nb_scolarises_producteur * proportion_dans_tranche
                        age_groups[age_range]['scolarises'] += scolarises_estimes
        
        # Arrondir les nombres de scolarisés
        for age_range in age_groups.keys():
            age_groups[age_range]['scolarises'] = round(age_groups[age_range]['scolarises'])
        
        # Calculer les taux de scolarisation par tranche d'âge
        scolarisation_par_age = []
        total_scolarises_estimes = 0
        
        for age_range, data in age_groups.items():
            total_enfants_tranche = data['total']
            scolarises = data['scolarises']
            non_scolarises = total_enfants_tranche - scolarises
            taux = round((scolarises / total_enfants_tranche * 100) if total_enfants_tranche > 0 else 0, 2)
            
            # Accumuler le nombre total d'enfants scolarisés estimés (3-18 ans)
            total_scolarises_estimes += scolarises
            
            scolarisation_par_age.append({
                'tranche_age': age_range,
                'total': total_enfants_tranche,
                'scolarises': scolarises,
                'non_scolarises': non_scolarises,
                'taux_scolarisation': taux
            })
    if not use_current_data:
        # Pour les données historiques, pas de détail par tranche d'âge
        impact_stats = {
            'with_kit': {'children_school_age': 0, 'schooled_estimated': 0},
            'without_kit': {'children_school_age': 0, 'schooled_estimated': 0}
        }
        scolarisation_par_age = []
        total_scolarises_estimes = enfants_scolarises
    
    # Statistiques adultes
    agregats_adultes = base_queryset.aggregate(
        total_adultes=Sum('nb_adultes_plus_18'),
        total_hommes=Sum('nb_hommes_adultes'),
        total_femmes=Sum('nb_femmes_adultes'),
        personnes_handicap=Count('id', filter=Q(personne_handicap_foyer=True))
    )
    
    total_adultes = agregats_adultes['total_adultes'] or 0
    total_hommes_adultes = agregats_adultes['total_hommes'] or 0
    total_femmes_adultes = agregats_adultes['total_femmes'] or 0
    personnes_handicap = agregats_adultes['personnes_handicap'] or 0
    
    # Statistiques hygiène
    hygiene_stats = {
        'poubelles_triees': base_queryset.filter(a_poubelles_triees=True).count(),
        'wc_maison': base_queryset.filter(wc_maison=True).count(),
        'eau_potable': base_queryset.filter(eau_potable=True).count(),
        'assurance_sante': base_queryset.filter(a_assurance_sante=True).count(),
    }
    
    # Statistiques environnement
    environnement_stats = {
        'respecte_dina': base_queryset.filter(respecte_dina=True).count(),
        'ne_brule_pas_foret': base_queryset.filter(ne_brule_pas_foret=True).count(),
        'ne_coupe_pas_foret': base_queryset.filter(ne_coupe_pas_foret=True).count(),
        'pratique_tavy': base_queryset.filter(pratique_tavy=True).count(),
    }
    
    # Statistiques activités
    # dotation, agr1, agr2 sont des champs texte (on compte les non-vides)
    # mahavelona est un booléen
    
    # Compter les activités AGR et extraire les types
    agr1_producteurs = base_queryset.exclude(Q(agr1='') | Q(agr1__isnull=True))
    agr2_producteurs = base_queryset.exclude(Q(agr2='') | Q(agr2__isnull=True))
    
    # Compter les types d'AGR différents
    agr1_types = {}
    for p in agr1_producteurs:
        agr_type = p.agr1.strip() if p.agr1 else ''
        if agr_type:
            agr1_types[agr_type] = agr1_types.get(agr_type, 0) + 1
    
    agr2_types = {}
    for p in agr2_producteurs:
        agr_type = p.agr2.strip() if p.agr2 else ''
        if agr_type:
            agr2_types[agr_type] = agr2_types.get(agr_type, 0) + 1
    
    # Compter les producteurs avec dotations (nouveau modèle Dotation)
    producteurs_avec_dotations = base_queryset.filter(dotations__isnull=False).distinct().count()
    
    # Compter les dotations par type
    from producteurs.models import Dotation
    dotation_types = {}
    # Agrégation DB des dotations par type au lieu de boucle Python
    dotation_aggregates = (
        Dotation.objects.filter(producteur__in=base_queryset)
        .values('type_dotation')
        .annotate(count=Count('id'))
    )
    dotation_types = {}
    for row in dotation_aggregates:
        type_label = dict(Dotation._meta.get_field('type_dotation').choices).get(row['type_dotation'], row['type_dotation'])
        dotation_types[type_label] = row['count']
    
    activites_stats = {
        'pratique_elevage': base_queryset.filter(pratique_elevage=True).count(),
        'pratique_peche': base_queryset.filter(pratique_peche=True).count(),
        'pratique_chasse': base_queryset.filter(pratique_chasse=True).count(),
        # Nouvelles activités
        'dotation': producteurs_avec_dotations,
        'agr1': agr1_producteurs.count(),
        'agr2': agr2_producteurs.count(),
        'mahavelona': base_queryset.filter(mahavelona=True).count(),
        # Pourcentages
        'dotation_pct': round((producteurs_avec_dotations / actifs * 100) if actifs > 0 else 0, 2),
        'agr1_pct': round((agr1_producteurs.count() / actifs * 100) if actifs > 0 else 0, 2),
        'agr2_pct': round((agr2_producteurs.count() / actifs * 100) if actifs > 0 else 0, 2),
        'mahavelona_pct': round((base_queryset.filter(mahavelona=True).count() / actifs * 100) if actifs > 0 else 0, 2),
        # Détails par type d'AGR
        'agr1_types': [{'type': k, 'count': v} for k, v in sorted(agr1_types.items(), key=lambda x: x[1], reverse=True)],
        'agr2_types': [{'type': k, 'count': v} for k, v in sorted(agr2_types.items(), key=lambda x: x[1], reverse=True)],
        # Détails par type de dotation
        'dotation_types': [{'type': k, 'count': v} for k, v in sorted(dotation_types.items(), key=lambda x: x[1], reverse=True)],
    }
    
    # Statistiques vie communautaire
    community_stats = {
        'participe_travaux_communautaires': base_queryset.filter(participe_travaux_communautaires=True).count(),
        'participe_protection_environnement': base_queryset.filter(participe_protection_environnement=True).count(),
        'respecte_dina': base_queryset.filter(respecte_dina=True).count(),
    }
    
    # Évolution mensuelle (12 derniers mois)
    date_debut = datetime.now() - timedelta(days=365)
    evolution_mensuelle = list(
        Producteur.objects.filter(date_adhesion_cooperative__gte=date_debut)
        .annotate(mois=TruncMonth('date_adhesion_cooperative'))
        .values('mois')
        .annotate(count=Count('id'))
        .order_by('mois')
    )
    
    # Répartition par tranches d'âge
    tranches_age = Producteur.objects.filter(actif=True, date_naissance__isnull=False).aggregate(
        moins_25=Count(Case(When(date_naissance__gte=datetime.now().date() - timedelta(days=25*365), then=1))),
        age_25_35=Count(Case(When(
            date_naissance__gte=datetime.now().date() - timedelta(days=35*365),
            date_naissance__lt=datetime.now().date() - timedelta(days=25*365),
            then=1
        ))),
        age_35_45=Count(Case(When(
            date_naissance__gte=datetime.now().date() - timedelta(days=45*365),
            date_naissance__lt=datetime.now().date() - timedelta(days=35*365),
            then=1
        ))),
        age_45_plus=Count(Case(When(date_naissance__lt=datetime.now().date() - timedelta(days=45*365), then=1)))
    )
    
    # Calcul Taux Impact
    def calc_rate(num, den):
        return round((num / den * 100), 2) if den > 0 else 0

    taux_avec_kit = calc_rate(impact_stats['with_kit']['schooled_estimated'], impact_stats['with_kit']['children_school_age'])
    taux_sans_kit = calc_rate(impact_stats['without_kit']['schooled_estimated'], impact_stats['without_kit']['children_school_age'])

    taux_actifs_filtres = round((actifs / total_filtres * 100), 2) if total_filtres > 0 else 0

    comparaison_annee = {
        'annee': selected_year,
        'annee_precedente': (selected_year - 1) if selected_year else None,
        'taux': round((total_scolarises_estimes / enfants_en_age_scolaire * 100), 2) if enfants_en_age_scolaire > 0 else 0,
        'taux_precedent': None,
        'delta': None,
    }
    if selected_year:
        try:
            from history.models import SocialIndicatorHistory
            prev_qs = SocialIndicatorHistory.objects.filter(
                annee=selected_year - 1,
                type_indicateur='scolarisation'
            )
            if villages and any(villages):
                prev_qs = prev_qs.filter(producteur__village__in=[v for v in villages if v])
            if communes and any(communes):
                prev_qs = prev_qs.filter(producteur__commune__in=[c for c in communes if c])
            if fokontanys and any(fokontanys):
                prev_qs = prev_qs.filter(producteur__fokontany__in=[f for f in fokontanys if f])

            prev_total = 0
            for record in prev_qs:
                if record.valeur_numerique is not None:
                    prev_total += float(record.valeur_numerique)
            if prev_qs.exists():
                comparaison_annee['taux_precedent'] = round((prev_total / max(total_enfants, 1)) * 100, 2)
                comparaison_annee['delta'] = round(
                    comparaison_annee['taux'] - comparaison_annee['taux_precedent'],
                    2
                )
        except Exception:
            pass

    dashboard_data = {
        'impact_dotation': {
            'taux_avec_kit': taux_avec_kit,
            'taux_sans_kit': taux_sans_kit,
            'enfants_concernes_avec_kit': impact_stats['with_kit']['children_school_age'],
            'enfants_concernes_sans_kit': impact_stats['without_kit']['children_school_age']
        },
        'global': {
            'total': total,  # Total de TOUS les producteurs (actifs + inactifs)
            'total_filtres': total_filtres,
            'actifs': actifs_sans_filtres,  # Nombre d'actifs SANS filtres (pour cohérence avec total)
            'actifs_filtres': actifs,  # Nombre d'actifs avec filtres appliqués
            'inactifs': inactifs,
            'inactifs_filtres': inactifs_filtres,
            'taux_actifs': round((actifs_sans_filtres / total * 100) if total > 0 else 0, 2),
            'taux_actifs_filtres': taux_actifs_filtres,
        },
        'demographics': {
            'par_genre': par_genre,
            'par_education': par_education,
            'par_statut': par_statut,
            'tranches_age': tranches_age,
            'age_distribution': age_distribution_list,
        },
        'geographic': {
            'top_communes': top_communes,
        },
        'cooperative': {
            'par_cooperative': par_cooperative,
            'responsables': responsables,
        },
        'leadership': {
            'femmes_leaders': femmes_leaders,
            'paysans_relais': paysans_relais,
            'satellite_floraison': base_queryset.filter(satellite_floraison=True).count(),
            'pourcentage_leaders': round((femmes_leaders / actifs * 100) if actifs > 0 else 0, 2),
        },
        'enfants': {
            'total': total_enfants,
            'scolarises': enfants_scolarises,
            'non_scolarises': enfants_non_scolarises,
            'en_age_scolaire': enfants_en_age_scolaire,
            'enfants_age_scolaire': enfants_en_age_scolaire,
            'enfants_scolarises': enfants_scolarises,
            'taux_scolarisation': round((enfants_scolarises / enfants_en_age_scolaire * 100) if enfants_en_age_scolaire > 0 else 0, 2),
            'methode_calcul': scolarisation_methode,
            'scolarisation_slots': scolarisation_slots,
            'comparaison_annee': comparaison_annee,
            'scolarisation_par_age': scolarisation_par_age,
        },
        'adultes': {
            'total': total_adultes,
            'hommes': total_hommes_adultes,
            'femmes': total_femmes_adultes,
            'personnes_handicap': personnes_handicap,
        },
        'hygiene': hygiene_stats,
        'environnement': environnement_stats,
        'activites': activites_stats,
        'community': community_stats,
        'evolution': evolution_mensuelle,
    }
    
    return Response(dashboard_data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_producteurs(request):
    """
    Statistiques spécifiques aux producteurs
    GET /api/dashboard/producteurs/
    """
    data = {
        'total': Producteur.objects.count(),
        'actifs': Producteur.objects.actifs().count(),
        'par_genre': list(
            Producteur.objects.filter(actif=True)
            .values('sexe')
            .annotate(count=Count('id'))
        ),
        'par_village': list(
            Producteur.objects.filter(actif=True)
            .values('village')
            .annotate(count=Count('id'))
            .order_by('-count')[:10]
        ),
    }
    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_hygiene(request):
    """
    Statistiques spécifiques à l'hygiène
    GET /api/dashboard/hygiene/
    """
    base_queryset = _dashboard_base_queryset(request)
    actifs = base_queryset.count()

    def _count(**filters):
        return base_queryset.filter(**filters).count()

    data = {
        'total_producteurs': actifs,
        'poubelles_triees': {
            'count': _count(a_poubelles_triees=True),
            'pourcentage': 0,
        },
        'wc_maison': {
            'count': _count(wc_maison=True),
            'pourcentage': 0,
        },
        'wc_champ': {
            'count': _count(wc_champ=True),
            'pourcentage': 0,
        },
        'eau_potable': {
            'count': _count(eau_potable=True),
            'pourcentage': 0,
        },
        'assurance_sante': {
            'count': _count(a_assurance_sante=True),
            'pourcentage': 0,
        },
    }

    # Calculer les pourcentages
    if actifs > 0:
        for key in ['poubelles_triees', 'wc_maison', 'wc_champ', 'eau_potable', 'assurance_sante']:
            data[key]['pourcentage'] = round((data[key]['count'] / actifs) * 100, 2)

    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_hygiene_absents(request):
    """
    Liste des producteurs ACTIFS manquant un indicateur d'hygiène/santé.
    GET /api/dashboard/hygiene/absents/?champ=wc_maison

    Champs autorisés (liste blanche) : wc_maison, wc_champ, eau_potable,
    a_assurance_sante, a_poubelles_triees.
    Retourne : [{ id, code, nom_complet, village, commune, fokontany, telephone }]
    """
    CHAMPS_AUTORISES = {
        'wc_maison': 'WC à la maison',
        'wc_champ': 'WC au champ',
        'eau_potable': "Accès à l'eau potable",
        'a_assurance_sante': 'Assurance santé',
        'a_poubelles_triees': 'Poubelles triées',
    }
    champ = request.query_params.get('champ', '')
    if champ not in CHAMPS_AUTORISES:
        return Response(
            {'detail': f"Champ invalide. Champs autorisés : {', '.join(CHAMPS_AUTORISES)}."},
            status=400,
        )

    producteurs = _dashboard_base_queryset(request).filter(**{champ: False}).order_by('code')
    data = [
        {
            'id': p.id,
            'code': p.code,
            'nom_complet': p.nom_complet,
            'village': p.village,
            'commune': p.commune,
            'fokontany': p.fokontany,
            'telephone': p.telephone,
        }
        for p in producteurs
    ]
    return Response({
        'champ': champ,
        'label': CHAMPS_AUTORISES[champ],
        'count': len(data),
        'producteurs': data,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_environnement(request):
    """
    Statistiques spécifiques à l'environnement
    GET /api/dashboard/environnement/
    """
    base_queryset = _dashboard_base_queryset(request)
    actifs = base_queryset.count()

    data = {
        'total_producteurs': actifs,
        'respecte_dina': base_queryset.filter(respecte_dina=True).count(),
        'ne_brule_pas_foret': base_queryset.filter(ne_brule_pas_foret=True).count(),
        'ne_coupe_pas_foret': base_queryset.filter(ne_coupe_pas_foret=True).count(),
        'pratique_tavy': base_queryset.filter(pratique_tavy=True).count(),
        'pratique_peche': base_queryset.filter(pratique_peche=True).count(),
    }

    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_villages_and_communes(request):
    """
    Retourne la liste des villages et communes disponibles
    GET /api/dashboard/villages-communes/
    """
    villages = list(
        Producteur.objects.filter(actif=True)
        .exclude(village__isnull=True)
        .exclude(village__exact='')
        .values_list('village', flat=True)
        .distinct()
        .order_by('village')
    )
    # Inclure les villages de référence
    villages_ref = list(
        VillageReference.objects.exclude(name__isnull=True).exclude(name__exact='')
        .values_list('name', flat=True).distinct().order_by('name')
    )
    villages = sorted(list({*villages, *villages_ref}))
    
    communes = list(
        Producteur.objects.filter(actif=True)
        .exclude(commune__isnull=True)
        .exclude(commune__exact='')
        .values_list('commune', flat=True)
        .distinct()
        .order_by('commune')
    )
    communes_ref = list(
        VillageReference.objects.exclude(commune__isnull=True).exclude(commune__exact='')
        .values_list('commune', flat=True).distinct().order_by('commune')
    )
    communes = sorted(list({*communes, *communes_ref}))
    
    fokontanys = list(
        Producteur.objects.filter(actif=True)
        .exclude(fokontany__isnull=True)
        .exclude(fokontany__exact='')
        .values_list('fokontany', flat=True)
        .distinct()
        .order_by('fokontany')
    )
    fokontanys_ref = list(
        VillageReference.objects.exclude(fokontany__isnull=True).exclude(fokontany__exact='')
        .values_list('fokontany', flat=True).distinct().order_by('fokontany')
    )
    fokontanys = sorted(list({*fokontanys, *fokontanys_ref}))
    
    return Response({
        'villages': villages,
        'communes': communes,
        'fokontanys': fokontanys
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def create_village(request):
    name = request.data.get('name')
    commune = request.data.get('commune')
    fokontany = request.data.get('fokontany')
    if not name or not str(name).strip():
        return Response({'detail': 'Le nom du village est requis.'}, status=400)
    obj, created = VillageReference.objects.get_or_create(
        name=str(name).strip(),
        commune=(str(commune).strip() if commune else None),
        fokontany=(str(fokontany).strip() if fokontany else None),
    )
    return Response({
        'id': obj.id,
        'name': obj.name,
        'commune': obj.commune,
        'fokontany': obj.fokontany,
        'created': created,
    }, status=201 if created else 200)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def import_villages_excel(request):
    file = request.FILES.get('file')
    if not file:
        return Response({'detail': 'Aucun fichier reçu.'}, status=400)
    try:
        import pandas as pd
        df = pd.read_excel(file)
    except Exception as e:
        return Response({'detail': f"Fichier invalide: {e}"}, status=400)

    # Normaliser noms de colonnes
    cols = {c.lower().strip(): c for c in df.columns if isinstance(c, str)}
    name_col = cols.get('village') or cols.get('name')
    commune_col = cols.get('commune')
    fokontany_col = cols.get('fokontany')
    if not name_col:
        return Response({'detail': 'Colonne "village" ou "name" manquante.'}, status=400)

    created_count = 0
    skipped = 0
    for _, row in df.iterrows():
        name = str(row.get(name_col) or '').strip()
        if not name:
            skipped += 1
            continue
        commune = row.get(commune_col) if commune_col else None
        fokontany = row.get(fokontany_col) if fokontany_col else None
        _, created = VillageReference.objects.get_or_create(
            name=name,
            commune=(str(commune).strip() if isinstance(commune, str) else commune),
            fokontany=(str(fokontany).strip() if isinstance(fokontany, str) else fokontany),
        )
        if created:
            created_count += 1
        else:
            skipped += 1

    return Response({
        'created': created_count,
        'skipped': skipped,
        'total_rows': int(df.shape[0])
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_enfants(request):
    """
    Statistiques spécifiques aux enfants
    GET /api/dashboard/enfants/
    """
    base_queryset = _dashboard_base_queryset(request)

    agregats = base_queryset.aggregate(
        total_garcons=Sum(F('nb_enfants_garcons') + F('nb_autres_garcons')),
        total_filles=Sum(F('nb_enfants_filles') + F('nb_autres_filles')),
    )

    total_garcons = agregats['total_garcons'] or 0
    total_filles = agregats['total_filles'] or 0
    total_enfants = total_garcons + total_filles

    # Scolarisation : cascade robuste (voir _children_schooling_stats).
    # Les champs agrégés (nb_enfants_scolarises) et les slots
    # annee_naissance_enfant_N sont lacunaires dans la base, d'où le recours
    # aux compteurs d'enfants déclarés (source fiable).
    annee_actuelle = datetime.now().year
    schooling = _children_schooling_stats(base_queryset, annee_actuelle)

    data = {
        'total_enfants': total_enfants,
        'enfants_scolarises': schooling['enfants_scolarises'],
        'enfants_non_scolarises': schooling['enfants_non_scolarises'],
        'enfants_en_age_scolaire': schooling['enfants_en_age_scolaire'],
        'total_garcons': total_garcons,
        'total_filles': total_filles,
        'taux_scolarisation': schooling['taux_scolarisation'],
        'methode_calcul': schooling['methode'],
        'slots': {
            'age': schooling['slots_age'],
            'scolarises': schooling['slots_scolarises'],
        },
    }

    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_scolarisation_evolution(request):
    """
    Évolution du taux de scolarisation par année (#12).

    GET /api/dashboard/scolarisation/evolution/

    Sources :
    - années archivées : `ProducteurSnapshot` (un enregistrement par producteur/an)
    - année courante   : calcul live (cascade `_children_schooling_stats`)

    Paramètres : village, commune, fokontany (multi-valeurs)
    """
    villages = [v for v in request.query_params.getlist('village') if v]
    communes = [c for c in request.query_params.getlist('commune') if c]
    fokontanys = [f for f in request.query_params.getlist('fokontany') if f]

    annee_courante = datetime.now().year
    annees = []

    try:
        from history.models import ProducteurSnapshot
    except ImportError:
        ProducteurSnapshot = None

    if ProducteurSnapshot is not None:
        snapshots = ProducteurSnapshot.objects.all()
        if villages:
            snapshots = snapshots.filter(village__in=villages)
        if communes:
            snapshots = snapshots.filter(commune__in=communes)
        if fokontanys:
            snapshots = snapshots.filter(fokontany__in=fokontanys)

        for row in (
            snapshots
            .values('annee')
            .annotate(
                total_enfants=Sum(F('nb_enfants_garcons') + F('nb_enfants_filles')),
                scolarises=Sum('nb_enfants_scolarises'),
                non_scolarises=Sum('nb_enfants_non_scolarises'),
                producteurs=Count('id'),
            )
            .order_by('annee')
        ):
            if row['annee'] == annee_courante:
                # L'année en cours est calculée en live (plus fiable que le snapshot)
                continue
            total = row['total_enfants'] or 0
            scolarises = row['scolarises'] or 0
            if total == 0 and scolarises == 0:
                continue
            annees.append({
                'annee': row['annee'],
                'source': 'snapshot',
                'total_enfants': total,
                'scolarises': scolarises,
                'non_scolarises': row['non_scolarises'] or 0,
                'taux_scolarisation': round((scolarises / total * 100) if total > 0 else 0, 2),
                'producteurs': row['producteurs'],
                'methode': 'snapshot_annuel',
            })

    # Année courante : calcul live
    live_qs = _dashboard_base_queryset(request)
    live = _children_schooling_stats(live_qs, annee_courante)
    annees.append({
        'annee': annee_courante,
        'source': 'live',
        'total_enfants': live['total_enfants'],
        'scolarises': live['enfants_scolarises'],
        'non_scolarises': live['enfants_non_scolarises'],
        'taux_scolarisation': live['taux_scolarisation'],
        'producteurs': live_qs.count(),
        'methode': live['methode'],
    })

    annees.sort(key=lambda a: a['annee'])

    total_cumule = sum(a['total_enfants'] for a in annees)
    scolarises_cumule = sum(a['scolarises'] for a in annees)

    return Response({
        'annees': annees,
        'nb_annees': len(annees),
        'annee_courante': annee_courante,
        'cumul': {
            'total_enfants': total_cumule,
            'scolarises': scolarises_cumule,
            'taux_scolarisation': round(
                (scolarises_cumule / total_cumule * 100) if total_cumule > 0 else 0, 2
            ),
        },
        'filtres_appliques': bool(villages or communes or fokontanys),
        'message_donnees': (
            "Une seule année de données est disponible : l'évolution annuelle "
            "s'affichera dès que d'autres années auront été archivées."
            if len(annees) <= 1 else None
        ),
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_production(request):
    """
    Estimations de production par culture avec filtres
    GET /api/dashboard/production/
    
    Paramètres de filtrage:
    - village: filtre par village(s) - peut être multiple
    - commune: filtre par commune(s) - peut être multiple
    - culture: filtre par type de culture (vanille, cafe, girofle, autre)
    """
    # Récupérer les paramètres de filtrage (support multi-sélection)
    villages = request.query_params.getlist('village', None)
    communes = request.query_params.getlist('commune', None)
    culture = request.query_params.getlist('culture', None)
    if culture:
        culture = [c for c in culture if c]
    if not culture:
        culture = None
    
    # Construire le queryset de base avec les parcelles actives
    base_queryset = Parcelle.objects.filter(active=True).select_related('producteur')
    
    # Appliquer les filtres géographiques
    if villages and any(villages):
        villages = [v for v in villages if v]
        if villages:
            base_queryset = base_queryset.filter(producteur__village__in=villages)
    
    if communes and any(communes):
        communes = [c for c in communes if c]
        if communes:
            base_queryset = base_queryset.filter(producteur__commune__in=communes)
    
    # Filtrage par culture : utiliser productions_par_culture JSONField (#16 : multi)
    if culture:
        culture_q = Q()
        for c in culture:
            culture_q |= Q(productions_par_culture__has_key=c)
        base_queryset = base_queryset.filter(culture_q)
    
    # Calculer les statistiques selon la/les culture(s) sélectionnée(s)
    if culture:
        # Agréger uniquement la production des cultures sélectionnées
        from django.db.models import F, Value
        from django.db.models.functions import Cast
        from django.db.models import FloatField

        # Extraire la production des cultures spécifiques depuis le JSON
        total_parcelles = base_queryset.count()
        total_superficie = base_queryset.aggregate(total=Sum('dimension_ha'))['total'] or 0

        # Calculer la production totale pour ces cultures (somme sur les cultures demandées)
        total_production = 0
        for parcelle in base_queryset:
            if parcelle.productions_par_culture and isinstance(parcelle.productions_par_culture, dict):
                for c in culture:
                    if c in parcelle.productions_par_culture:
                        try:
                            total_production += float(parcelle.productions_par_culture[c])
                        except (ValueError, TypeError):
                            pass

        agregats_production = {
            'total_parcelles': total_parcelles,
            'total_superficie': total_superficie,
            'total_production_estimee': total_production,
            'total_pieds_vanille': base_queryset.aggregate(total=Sum('nombre_pieds'))['total'] or 0 if 'vanille' in culture else 0,
        }

        # Production par culture (une entrée par culture demandée).
        # #15 — La superficie est celle des parcelles où la culture est
        # réellement pratiquée (et non la superficie globale répétée), et
        # `nb_parcelles` inclut aussi les parcelles dont c'est la culture
        # principale même sans production renseignée.
        from django.db.models import Q as _Q
        par_culture = []
        for c in culture:
            prod_c = 0
            superficie_c = 0.0
            for parcelle in base_queryset:
                productions = parcelle.productions_par_culture if isinstance(
                    parcelle.productions_par_culture, dict) else {}
                concerne = parcelle.culture_principale == c or c in productions
                if not concerne:
                    continue
                superficie_c += float(parcelle.dimension_ha or 0)
                if c in productions:
                    try:
                        prod_c += float(productions[c])
                    except (ValueError, TypeError):
                        pass
            q_c = _Q(productions_par_culture__has_key=c) | _Q(culture_principale=c)
            par_culture.append({
                'culture_principale': c,
                'nb_parcelles': base_queryset.filter(q_c).count(),
                'superficie_totale': round(superficie_c, 4),
                'production_estimee': prod_c
            })
    else:
        # Sans filtre culture : statistiques globales normales
        agregats_production = base_queryset.aggregate(
            total_parcelles=Count('id'),
            total_superficie=Sum('dimension_ha'),
            total_production_estimee=Sum('estimation_production_kg'),
            total_pieds_vanille=Sum('nombre_pieds'),
        )
        
        # Agréger toutes les productions par culture depuis productions_par_culture
        productions_par_culture_dict = {}
        for parcelle in base_queryset:
            if parcelle.productions_par_culture and isinstance(parcelle.productions_par_culture, dict):
                for cult, prod in parcelle.productions_par_culture.items():
                    if cult and prod:
                        cult_key = str(cult).strip().lower()
                        if cult_key not in productions_par_culture_dict:
                            productions_par_culture_dict[cult_key] = {
                                'nb_parcelles': 0,
                                'superficie_totale': 0,
                                'production_estimee': 0
                            }
                        try:
                            productions_par_culture_dict[cult_key]['nb_parcelles'] += 1
                            productions_par_culture_dict[cult_key]['production_estimee'] += float(prod)
                            productions_par_culture_dict[cult_key]['superficie_totale'] += float(parcelle.dimension_ha or 0)
                        except (ValueError, TypeError):
                            pass
        
        par_culture = [
            {
                'culture_principale': cult,
                'nb_parcelles': data['nb_parcelles'],
                'superficie_totale': data['superficie_totale'],
                'production_estimee': data['production_estimee']
            }
            for cult, data in productions_par_culture_dict.items()
        ]
        par_culture.sort(key=lambda x: x['production_estimee'], reverse=True)
    
    # Production par village, commune et producteur
    if culture:
        # Calculer les stats pour la culture sélectionnée
        villages_dict = {}
        communes_dict = {}
        producteurs_dict = {}
        
        for parcelle in base_queryset:
            if not (parcelle.productions_par_culture and isinstance(parcelle.productions_par_culture, dict)):
                continue
            prod = 0.0
            found = False
            for c in culture:
                if c in parcelle.productions_par_culture:
                    try:
                        prod += float(parcelle.productions_par_culture[c])
                        found = True
                    except (ValueError, TypeError):
                        pass
            if not found:
                continue
                
                # Par village
                village = parcelle.producteur.village
                if village not in villages_dict:
                    villages_dict[village] = {'nb_parcelles': 0, 'superficie_totale': 0, 'production_estimee': 0}
                villages_dict[village]['nb_parcelles'] += 1
                villages_dict[village]['superficie_totale'] += float(parcelle.dimension_ha or 0)
                villages_dict[village]['production_estimee'] += prod
                
                # Par commune
                commune = parcelle.producteur.commune
                if commune not in communes_dict:
                    communes_dict[commune] = {'nb_parcelles': 0, 'superficie_totale': 0, 'production_estimee': 0}
                communes_dict[commune]['nb_parcelles'] += 1
                communes_dict[commune]['superficie_totale'] += float(parcelle.dimension_ha or 0)
                communes_dict[commune]['production_estimee'] += prod
                
                # Par producteur
                prod_key = (parcelle.producteur.code, parcelle.producteur.nom, parcelle.producteur.prenom)
                if prod_key not in producteurs_dict:
                    producteurs_dict[prod_key] = {'nb_parcelles': 0, 'superficie_totale': 0, 'production_estimee': 0}
                producteurs_dict[prod_key]['nb_parcelles'] += 1
                producteurs_dict[prod_key]['superficie_totale'] += float(parcelle.dimension_ha or 0)
                producteurs_dict[prod_key]['production_estimee'] += prod
        
        par_village = [{'producteur__village': v, **data} for v, data in villages_dict.items()]
        par_village.sort(key=lambda x: x['production_estimee'], reverse=True)
        par_village = par_village[:10]
        
        par_commune = [{'producteur__commune': c, **data} for c, data in communes_dict.items()]
        par_commune.sort(key=lambda x: x['production_estimee'], reverse=True)
        par_commune = par_commune[:10]
        
        top_producteurs = [
            {
                'producteur__code': k[0],
                'producteur__nom': k[1],
                'producteur__prenom': k[2],
                **data
            } 
            for k, data in producteurs_dict.items()
        ]
        top_producteurs.sort(key=lambda x: x['production_estimee'], reverse=True)
        top_producteurs = top_producteurs[:10]
        
        # Par type de vanille (seulement si vanille fait partie des cultures #16)
        par_type_vanille = []
        if 'vanille' in culture:
            type_vanille_dict = {}
            for parcelle in base_queryset:
                if parcelle.type_vanille and parcelle.productions_par_culture and 'vanille' in parcelle.productions_par_culture:
                    try:
                        prod = float(parcelle.productions_par_culture['vanille'])
                    except (ValueError, TypeError):
                        continue
                    
                    tv = parcelle.type_vanille
                    if tv not in type_vanille_dict:
                        type_vanille_dict[tv] = {'nb_parcelles': 0, 'superficie_totale': 0, 'production_estimee': 0, 'nb_pieds': 0}
                    type_vanille_dict[tv]['nb_parcelles'] += 1
                    type_vanille_dict[tv]['superficie_totale'] += float(parcelle.dimension_ha or 0)
                    type_vanille_dict[tv]['production_estimee'] += prod
                    type_vanille_dict[tv]['nb_pieds'] += parcelle.nombre_pieds or 0
            
            par_type_vanille = [{'type_vanille': tv, **data} for tv, data in type_vanille_dict.items()]
            par_type_vanille.sort(key=lambda x: x['production_estimee'], reverse=True)
        
        # Par certification (pour la culture sélectionnée)
        certification_dict = {}
        for parcelle in base_queryset.filter(certifiee=True):
            if not (parcelle.productions_par_culture and isinstance(parcelle.productions_par_culture, dict)):
                continue
            prod = 0.0
            found = False
            for c in culture:
                if c in parcelle.productions_par_culture:
                    try:
                        prod += float(parcelle.productions_par_culture[c])
                        found = True
                    except (ValueError, TypeError):
                        pass
            if not found:
                continue
                
                cert = parcelle.type_certification
                if cert not in certification_dict:
                    certification_dict[cert] = {'nb_parcelles': 0, 'superficie_totale': 0, 'production_estimee': 0}
                certification_dict[cert]['nb_parcelles'] += 1
                certification_dict[cert]['superficie_totale'] += float(parcelle.dimension_ha or 0)
                certification_dict[cert]['production_estimee'] += prod
        
        par_certification = [{'type_certification': cert, **data} for cert, data in certification_dict.items()]
        par_certification.sort(key=lambda x: x['production_estimee'], reverse=True)
    else:
        # Sans filtre culture : utiliser les agrégations normales
        par_village = list(
            base_queryset
            .values('producteur__village')
            .annotate(
                nb_parcelles=Count('id'),
                superficie_totale=Sum('dimension_ha'),
                production_estimee=Sum('estimation_production_kg')
            )
            .order_by('-production_estimee')[:10]
        )
        
        par_commune = list(
            base_queryset
            .values('producteur__commune')
            .annotate(
                nb_parcelles=Count('id'),
                superficie_totale=Sum('dimension_ha'),
                production_estimee=Sum('estimation_production_kg')
            )
            .order_by('-production_estimee')[:10]
        )
        
        par_type_vanille = list(
            base_queryset.filter(type_vanille__isnull=False)
            .values('type_vanille')
            .annotate(
                nb_parcelles=Count('id'),
                superficie_totale=Sum('dimension_ha'),
                production_estimee=Sum('estimation_production_kg'),
                nb_pieds=Sum('nombre_pieds')
            )
            .order_by('-production_estimee')
        )
        
        par_certification = list(
            base_queryset.filter(certifiee=True)
            .values('type_certification')
            .annotate(
                nb_parcelles=Count('id'),
                superficie_totale=Sum('dimension_ha'),
                production_estimee=Sum('estimation_production_kg')
            )
            .order_by('-production_estimee')
        )
        
        top_producteurs = list(
            base_queryset
            .values('producteur__code', 'producteur__nom', 'producteur__prenom')
            .annotate(
                nb_parcelles=Count('id'),
                superficie_totale=Sum('dimension_ha'),
                production_estimee=Sum('estimation_production_kg')
            )
            .order_by('-production_estimee')[:10]
        )
    
    # ===== #36 — Réalisations réelles + résilience (vendu / consommé) =====
    from history.models import AGRHistory, ProductionHistory

    # Réalisations réelles par année × culture (ProductionHistory) —
    # volontairement NON cumulé : une ligne par année, aucune somme inter-années.
    geo_parcelles = Parcelle.objects.filter(active=True)
    if villages and any(villages):
        geo_parcelles = geo_parcelles.filter(producteur__village__in=[v for v in villages if v])
    if communes and any(communes):
        geo_parcelles = geo_parcelles.filter(producteur__commune__in=[c for c in communes if c])
    realisations_qs = ProductionHistory.objects.filter(parcelle__in=geo_parcelles)
    if culture:
        realisations_qs = realisations_qs.filter(culture=culture)
    realisations = [
        {
            'annee': r['annee'],
            'culture': r['culture'],
            'produite_kg': float(r['produite'] or 0),
            'revenu_ar': float(r['revenu'] or 0),
        }
        for r in realisations_qs.values('annee', 'culture')
        .annotate(produite=Sum('quantite_kg'), revenu=Sum('revenu_total'))
        .order_by('-annee', 'culture')
    ]

    # Résilience « autres produits » (AGR) : quelle part est vendue / consommée ?
    # Ventilé par année et par type — non cumulé, pourcentages calculés par ligne.
    resilience_qs = AGRHistory.objects.all()
    if villages and any(villages):
        resilience_qs = resilience_qs.filter(producteur__village__in=[v for v in villages if v])
    if communes and any(communes):
        resilience_qs = resilience_qs.filter(producteur__commune__in=[c for c in communes if c])
    resilience_agr = []
    for r in resilience_qs.values('annee', 'type_agr').annotate(
            produite=Sum('quantite_produite'),
            vendue=Sum('quantite_vendue'),
            consommee=Sum('quantite_consommee'),
    ).order_by('-annee', 'type_agr'):
        produite = float(r['produite'] or 0)
        vendue = float(r['vendue'] or 0)
        consommee = float(r['consommee'] or 0)
        # Dénominateur : production si saisie, sinon ventes + consommation
        # (la production n'est pas toujours renseignée dans les imports).
        denom = produite if produite > 0 else (vendue + consommee)
        resilience_agr.append({
            'annee': r['annee'],
            'type_agr': r['type_agr'],
            'produite': produite,
            'vendue': vendue,
            'consommee': consommee,
            'pct_vendu': round(vendue / denom * 100, 1) if denom > 0 else None,
            'pct_consomme': round(consommee / denom * 100, 1) if denom > 0 else None,
        })

    # Calculer les rendements moyens
    total_superficie = float(agregats_production['total_superficie'] or 0)
    total_production = float(agregats_production['total_production_estimee'] or 0)
    rendement_moyen = round(total_production / total_superficie, 2) if total_superficie > 0 else 0
    
    data = {
        'global': {
            'total_parcelles': agregats_production['total_parcelles'] or 0,
            'total_superficie_ha': round(total_superficie, 2),
            'total_production_estimee_kg': round(total_production, 2),
            'total_pieds_vanille': agregats_production['total_pieds_vanille'] or 0,
            'rendement_moyen_kg_ha': rendement_moyen,
        },
        'par_village': par_village,
        'par_commune': par_commune,
        'par_type_vanille': par_type_vanille,
        'par_certification': par_certification,
        'top_producteurs': top_producteurs,
        'realisations': realisations,
        'resilience_agr': resilience_agr,
        'filtres_actifs': {
            'villages': villages if villages and any(villages) else [],
            'communes': communes if communes and any(communes) else [],
            'culture': culture,
        }
    }
    
    # par_culture : détail complet sans filtre, limité aux cultures choisies sinon (#16)
    data['par_culture'] = par_culture
    
    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_production_par_culture(request):
    """
    Détail des productions par type de culture (utilisant productions_par_culture JSONField)
    GET /api/dashboard/production-par-culture/
    
    Agrège toutes les productions depuis le champ productions_par_culture de chaque parcelle.
    Exemple: une parcelle avec {"vanille": 150, "cafe": 200} sera comptée pour les deux cultures.
    
    Paramètres de filtrage:
    - village: filtre par village(s) - peut être multiple
    - commune: filtre par commune(s) - peut être multiple
    """
    # Récupérer les paramètres de filtrage (support multi-sélection)
    villages = request.query_params.getlist('village', None)
    communes = request.query_params.getlist('commune', None)
    
    # Construire le queryset de base avec les parcelles actives
    base_queryset = Parcelle.objects.filter(active=True).select_related('producteur')
    
    # Appliquer les filtres
    if villages and any(villages):
        villages = [v for v in villages if v]
        if villages:
            base_queryset = base_queryset.filter(producteur__village__in=villages)
    
    if communes and any(communes):
        communes = [c for c in communes if c]
        if communes:
            base_queryset = base_queryset.filter(producteur__commune__in=communes)
    
    # Agréger les productions par culture
    productions_aggregees = {}
    nb_parcelles_par_culture = {}
    superficie_par_culture = {}

    for parcelle in base_queryset:
        if parcelle.productions_par_culture and isinstance(parcelle.productions_par_culture, dict):
            # #15 — la superficie de la parcelle compte pour chacune de ses cultures
            superficie_parcelle = float(parcelle.dimension_ha or 0)
            for culture, production in parcelle.productions_par_culture.items():
                if production and culture:
                    culture_key = str(culture).strip().lower()
                    if culture_key not in productions_aggregees:
                        productions_aggregees[culture_key] = 0
                        nb_parcelles_par_culture[culture_key] = 0
                        superficie_par_culture[culture_key] = 0.0

                    try:
                        productions_aggregees[culture_key] += float(production)
                        nb_parcelles_par_culture[culture_key] += 1
                        superficie_par_culture[culture_key] += superficie_parcelle
                    except (ValueError, TypeError):
                        pass
    
    # Formater les résultats
    productions_detaillees = [
        {
            'culture': culture,
            'production_totale_kg': round(production, 2),
            'nb_parcelles': nb_parcelles_par_culture.get(culture, 0),
            'superficie_totale_ha': round(superficie_par_culture.get(culture, 0.0), 4),
            'production_moyenne_kg': round(production / nb_parcelles_par_culture.get(culture, 1), 2) if nb_parcelles_par_culture.get(culture, 0) > 0 else 0
        }
        for culture, production in productions_aggregees.items()
    ]
    
    # Trier par production totale décroissante
    productions_detaillees.sort(key=lambda x: x['production_totale_kg'], reverse=True)
    
    # Statistiques globales
    total_production = sum(p['production_totale_kg'] for p in productions_detaillees)
    
    data = {
        'productions_par_culture': productions_detaillees,
        'total_production_kg': round(total_production, 2),
        'nb_cultures_differentes': len(productions_detaillees),
        'filtres_actifs': {
            'villages': villages if villages and any(villages) else [],
            'communes': communes if communes and any(communes) else [],
        }
    }
    
    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_village_references(request):
    """
    Retourne la liste complète des villages de référence avec leurs informations
    GET /api/dashboard/village-references/
    
    Paramètres optionnels:
    - search: recherche dans le nom, commune, fokontany ou région
    - commune: filtre par commune
    - region: filtre par région
    """
    from .serializers import VillageReferenceSerializer
    
    queryset = VillageReference.objects.all()
    
    # Filtres
    search = request.query_params.get('search', None)
    if search:
        from django.db.models import Q
        queryset = queryset.filter(
            Q(name__icontains=search) |
            Q(commune__icontains=search) |
            Q(fokontany__icontains=search) |
            Q(region__icontains=search)
        )
    
    commune = request.query_params.get('commune', None)
    if commune:
        queryset = queryset.filter(commune=commune)
    
    region = request.query_params.get('region', None)
    if region:
        queryset = queryset.filter(region=region)
    
    serializer = VillageReferenceSerializer(queryset, many=True)
    return Response(serializer.data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_decisionnel(request):
    """
    Dashboard décisionnel multi-vues (commercial, impact, traçabilité, finance)
    GET /api/dashboard/decisionnel/

    Filtres:
    - village (multiple)
    - commune (multiple)
    - certification (multiple)
    - date_from (YYYY-MM-DD)
    - date_to (YYYY-MM-DD)
    """
    villages = [v for v in request.query_params.getlist('village', []) if v]
    communes = [c for c in request.query_params.getlist('commune', []) if c]
    certifications = [c for c in request.query_params.getlist('certification', []) if c]
    date_from = request.query_params.get('date_from')
    date_to = request.query_params.get('date_to')

    start_date = None
    end_date = None
    try:
        if date_from:
            start_date = datetime.strptime(date_from, '%Y-%m-%d').date()
        if date_to:
            end_date = datetime.strptime(date_to, '%Y-%m-%d').date()
    except ValueError:
        return Response(
            {'detail': "Format de date invalide. Utiliser YYYY-MM-DD."},
            status=400
        )

    producteurs_qs = Producteur.objects.filter(actif=True)
    parcelles_qs = Parcelle.objects.filter(active=True, producteur__actif=True)
    agr_qs = AGR.objects.filter(active=True, producteur__actif=True)
    dotations_qs = Dotation.objects.filter(producteur__actif=True)
    bons_qs = BonCollecte.objects.all()
    lots_qs = LotTraitement.objects.all()
    colis_qs = Colis.objects.all()
    commandes_qs = CommandeExport.objects.all()
    chains_qs = TracabiliteChain.objects.all()

    if villages:
        producteurs_qs = producteurs_qs.filter(village__in=villages)
        parcelles_qs = parcelles_qs.filter(producteur__village__in=villages)
        agr_qs = agr_qs.filter(producteur__village__in=villages)
        dotations_qs = dotations_qs.filter(producteur__village__in=villages)
        bons_qs = bons_qs.filter(Q(village_marche__in=villages) | Q(producteur__village__in=villages))
        lots_qs = lots_qs.filter(bons_transport__fiche_collecte__bons_collecte__producteur__village__in=villages).distinct()
        colis_qs = colis_qs.filter(lot_traitement__bons_transport__fiche_collecte__bons_collecte__producteur__village__in=villages).distinct()
        commandes_qs = commandes_qs.filter(colis__lot_traitement__bons_transport__fiche_collecte__bons_collecte__producteur__village__in=villages).distinct()
        chains_qs = chains_qs.filter(producteur__village__in=villages)

    if communes:
        producteurs_qs = producteurs_qs.filter(commune__in=communes)
        parcelles_qs = parcelles_qs.filter(producteur__commune__in=communes)
        agr_qs = agr_qs.filter(producteur__commune__in=communes)
        dotations_qs = dotations_qs.filter(producteur__commune__in=communes)
        bons_qs = bons_qs.filter(Q(commune__in=communes) | Q(producteur__commune__in=communes))
        lots_qs = lots_qs.filter(bons_transport__fiche_collecte__bons_collecte__producteur__commune__in=communes).distinct()
        colis_qs = colis_qs.filter(lot_traitement__bons_transport__fiche_collecte__bons_collecte__producteur__commune__in=communes).distinct()
        commandes_qs = commandes_qs.filter(colis__lot_traitement__bons_transport__fiche_collecte__bons_collecte__producteur__commune__in=communes).distinct()
        chains_qs = chains_qs.filter(producteur__commune__in=communes)

    if certifications:
        cert_q = Q()
        for cert in certifications:
            cert_q |= Q(certification__iexact=cert) | Q(certification__icontains=cert)
        bons_qs = bons_qs.filter(cert_q)

    if start_date:
        bons_qs = bons_qs.filter(date_marche__gte=start_date)
        lots_qs = lots_qs.filter(date_debut__gte=start_date)
        colis_qs = colis_qs.filter(date_conditionnement__gte=start_date)
        commandes_qs = commandes_qs.filter(date_commande__gte=start_date)
        chains_qs = chains_qs.filter(date_creation__date__gte=start_date)
    if end_date:
        bons_qs = bons_qs.filter(date_marche__lte=end_date)
        lots_qs = lots_qs.filter(date_debut__lte=end_date)
        colis_qs = colis_qs.filter(date_conditionnement__lte=end_date)
        commandes_qs = commandes_qs.filter(date_commande__lte=end_date)
        chains_qs = chains_qs.filter(date_creation__date__lte=end_date)

    if start_date or end_date:
        start_year = start_date.year if start_date else None
        end_year = end_date.year if end_date else None
        if start_year and end_year:
            dotations_qs = dotations_qs.filter(annee__gte=start_year, annee__lte=end_year)
        elif start_year:
            dotations_qs = dotations_qs.filter(annee__gte=start_year)
        elif end_year:
            dotations_qs = dotations_qs.filter(annee__lte=end_year)

    total_estimation_kg = float(parcelles_qs.aggregate(total=Sum('estimation_production_kg'))['total'] or 0)
    total_collecte_kg = float(bons_qs.aggregate(total=Sum('poids_accepte'))['total'] or 0)
    total_achat_ar = float(bons_qs.aggregate(total=Sum('montant_total_achat'))['total'] or 0)
    total_export_usd = float(commandes_qs.aggregate(total=Sum('valeur_commande'))['total'] or 0)
    couverture_collecte = round((total_collecte_kg / total_estimation_kg * 100), 2) if total_estimation_kg > 0 else 0
    ecart_collecte = round(total_estimation_kg - total_collecte_kg, 2)
    prix_moyen = round((total_achat_ar / total_collecte_kg), 2) if total_collecte_kg > 0 else 0

    producteurs_actifs = producteurs_qs.count()
    femmes_leaders = producteurs_qs.filter(femme_leader=True).count()
    taux_femmes_leaders = round((femmes_leaders / producteurs_actifs * 100), 2) if producteurs_actifs > 0 else 0
    producteurs_agr = agr_qs.values('producteur_id').distinct().count()
    revenu_agr_total = float(agr_qs.aggregate(total=Sum('revenu_annuel_estime'))['total'] or 0)
    producteurs_mahavelona = producteurs_qs.filter(mahavelona=True).count()
    kits_qs = dotations_qs.filter(type_dotation='kit_scolaire')
    producteurs_avec_kit = kits_qs.values('producteur_id').distinct().count()
    kits_total = int(kits_qs.aggregate(total=Sum('quantite'))['total'] or 0)
    couverture_kits = round((producteurs_avec_kit / producteurs_actifs * 100), 2) if producteurs_actifs > 0 else 0

    total_chains = chains_qs.count()
    complete_chains = chains_qs.exclude(bon_collecte_id__isnull=True).exclude(fiche_collecte_id__isnull=True).exclude(
        bon_transport_id__isnull=True
    ).exclude(lot_traitement_id__isnull=True).exclude(colis_id__isnull=True).exclude(commande_export_id__isnull=True).count()
    traceability_completion = round((complete_chains / total_chains * 100), 2) if total_chains > 0 else 0
    bons_anomalies = bons_qs.filter(poids_retour__gt=0).count()

    total_premium_ar = float(bons_qs.aggregate(total=Sum('montant_premium'))['total'] or 0)
    total_remboursements_ar = float(
        bons_qs.aggregate(total=Sum(F('remboursement_par_vanille') + F('remboursement_especes')))['total'] or 0
    )
    solde_avances_ar = float(bons_qs.aggregate(total=Sum('solde_avances'))['total'] or 0)

    by_certification = list(
        bons_qs.values('certification').annotate(
            nombre_collectes=Count('id'),
            poids_kg=Sum('poids_accepte'),
            montant_ar=Sum('montant_total_achat')
        ).order_by('-poids_kg')
    )
    by_paiement = list(
        bons_qs.values('mode_paiement').annotate(
            nombre=Count('id'),
            montant_ar=Sum('montant_total_achat')
        ).order_by('-montant_ar')
    )

    collecte_mensuelle = list(
        bons_qs.annotate(mois=TruncMonth('date_marche'))
        .values('mois')
        .annotate(
            collecte_kg=Sum('poids_accepte'),
            achats_ar=Sum('montant_total_achat'),
            anomalies=Count('id', filter=Q(poids_retour__gt=0))
        )
        .order_by('mois')
    )
    export_mensuel = list(
        commandes_qs.annotate(mois=TruncMonth('date_commande'))
        .values('mois')
        .annotate(exports_usd=Sum('valeur_commande'))
        .order_by('mois')
    )
    export_by_month = {item['mois'].strftime('%Y-%m'): float(item['exports_usd'] or 0) for item in export_mensuel if item['mois']}
    mois_count = len(collecte_mensuelle) if collecte_mensuelle else 1
    objectif_mensuel = (total_estimation_kg / mois_count) if mois_count > 0 else 0
    comparatif_mensuel = []
    for item in collecte_mensuelle:
        if not item['mois']:
            continue
        month_key = item['mois'].strftime('%Y-%m')
        collecte_kg = float(item['collecte_kg'] or 0)
        comparatif_mensuel.append({
            'mois': month_key,
            'objectif_kg': round(objectif_mensuel, 2),
            'collecte_kg': round(collecte_kg, 2),
            'ecart_kg': round(objectif_mensuel - collecte_kg, 2),
            'achats_ar': round(float(item['achats_ar'] or 0), 2),
            'exports_usd': round(export_by_month.get(month_key, 0), 2),
            'anomalies': item['anomalies'] or 0,
        })

    top_anomalies = list(
        bons_qs.filter(poids_retour__gt=0).annotate(
            perte_kg=F('poids_total_livre') - F('poids_accepte')
        ).values(
            'id',
            'numero_fabc',
            'date_marche',
            'producteur__code',
            'producteur__nom',
            'producteur__prenom',
            'commune',
            'village_marche',
            'certification',
            'poids_total_livre',
            'poids_accepte',
            'poids_retour',
            'perte_kg',
        ).order_by('-poids_retour')[:20]
    )
    for item in top_anomalies:
        total_livre = float(item.get('poids_total_livre') or 0)
        poids_retour = float(item.get('poids_retour') or 0)
        item['poids_total_livre'] = round(total_livre, 3)
        item['poids_accepte'] = round(float(item.get('poids_accepte') or 0), 3)
        item['poids_retour'] = round(poids_retour, 3)
        item['perte_kg'] = round(float(item.get('perte_kg') or 0), 3)
        item['taux_perte_pct'] = round((poids_retour / total_livre * 100), 2) if total_livre > 0 else 0
        if item.get('date_marche'):
            item['date_marche'] = item['date_marche'].isoformat()

    alerts = []
    if couverture_collecte < 60:
        alerts.append("Couverture collecte faible (< 60% de l'estimation).")
    if bons_anomalies > 0:
        alerts.append(f"{bons_anomalies} collecte(s) avec retour/anomalie de poids.")
    if traceability_completion < 85:
        alerts.append("Complétude de traçabilité < 85%.")

    return Response({
        'commercial': {
            'estimation_production_kg': round(total_estimation_kg, 2),
            'collecte_reelle_kg': round(total_collecte_kg, 2),
            'ecart_estimation_collecte_kg': ecart_collecte,
            'couverture_collecte_pct': couverture_collecte,
            'valeur_achats_ar': round(total_achat_ar, 2),
            'valeur_exports_usd': round(total_export_usd, 2),
            'prix_moyen_achat_ar_kg': prix_moyen,
            'par_certification': by_certification,
            'comparatif_mensuel': comparatif_mensuel,
        },
        'impact': {
            'producteurs_actifs': producteurs_actifs,
            'femmes_leaders': femmes_leaders,
            'taux_femmes_leaders_pct': taux_femmes_leaders,
            'producteurs_avec_agr': producteurs_agr,
            'revenu_agr_total_ar': round(revenu_agr_total, 2),
            'producteurs_mahavelona': producteurs_mahavelona,
            'producteurs_avec_kit': producteurs_avec_kit,
            'kits_distribues': kits_total,
            'couverture_kits_pct': couverture_kits,
        },
        'tracabilite': {
            'total_collectes': bons_qs.count(),
            'total_lots_traitement': lots_qs.count(),
            'total_colis': colis_qs.count(),
            'total_commandes_export': commandes_qs.count(),
            'total_chaines': total_chains,
            'chaines_completes': complete_chains,
            'taux_completude_pct': traceability_completion,
            'collectes_anomalies_poids': bons_anomalies,
            'top_anomalies': top_anomalies,
        },
        'finance': {
            'depenses_achats_ar': round(total_achat_ar, 2),
            'montant_premium_ar': round(total_premium_ar, 2),
            'remboursements_avances_ar': round(total_remboursements_ar, 2),
            'solde_avances_ar': round(solde_avances_ar, 2),
            'recettes_exports_usd': round(total_export_usd, 2),
            'par_mode_paiement': by_paiement,
        },
        'alerts': alerts,
        'filtres_actifs': {
            'villages': villages,
            'communes': communes,
            'certifications': certifications,
            'date_from': date_from,
            'date_to': date_to,
        }
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def dashboard_decisionnel_export(request):
    """
    Export du dashboard décisionnel.
    GET /api/dashboard/decisionnel/export/?format=excel|pdf
    """
    export_format = (request.query_params.get('format') or 'excel').lower()
    decision_response = dashboard_decisionnel(request)
    if decision_response.status_code != 200:
        return decision_response
    data = decision_response.data

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

    if export_format == 'excel':
        try:
            import openpyxl
            from openpyxl.styles import Font
        except ImportError:
            return Response(
                {'detail': "Le module openpyxl n'est pas installé sur le serveur."},
                status=501
            )

        wb = openpyxl.Workbook()
        ws_resume = wb.active
        ws_resume.title = "Resume"

        rows = [
            ("Section", "Indicateur", "Valeur"),
            ("Commercial", "Estimation production (kg)", data['commercial']['estimation_production_kg']),
            ("Commercial", "Collecte réelle (kg)", data['commercial']['collecte_reelle_kg']),
            ("Commercial", "Couverture collecte (%)", data['commercial']['couverture_collecte_pct']),
            ("Commercial", "Valeur achats (Ar)", data['commercial']['valeur_achats_ar']),
            ("Commercial", "Valeur exports (USD)", data['commercial']['valeur_exports_usd']),
            ("Impact", "Producteurs actifs", data['impact']['producteurs_actifs']),
            ("Impact", "Femmes leaders (%)", data['impact']['taux_femmes_leaders_pct']),
            ("Impact", "Producteurs avec AGR", data['impact']['producteurs_avec_agr']),
            ("Impact", "Couverture kits (%)", data['impact']['couverture_kits_pct']),
            ("Traçabilité", "Taux complétude (%)", data['tracabilite']['taux_completude_pct']),
            ("Traçabilité", "Anomalies de poids", data['tracabilite']['collectes_anomalies_poids']),
            ("Finance", "Dépenses achats (Ar)", data['finance']['depenses_achats_ar']),
            ("Finance", "Premium (Ar)", data['finance']['montant_premium_ar']),
            ("Finance", "Recettes export (USD)", data['finance']['recettes_exports_usd']),
        ]
        for ridx, row in enumerate(rows, start=1):
            for cidx, value in enumerate(row, start=1):
                ws_resume.cell(row=ridx, column=cidx, value=value)
        for cell in ws_resume[1]:
            cell.font = Font(bold=True)

        ws_comp = wb.create_sheet("Comparatif Mensuel")
        ws_comp.append(["Mois", "Objectif (kg)", "Collecte (kg)", "Ecart (kg)", "Achats (Ar)", "Exports (USD)", "Anomalies"])
        for item in data['commercial'].get('comparatif_mensuel', []):
            ws_comp.append([
                item['mois'], item['objectif_kg'], item['collecte_kg'], item['ecart_kg'],
                item['achats_ar'], item['exports_usd'], item['anomalies']
            ])
        for cell in ws_comp[1]:
            cell.font = Font(bold=True)

        ws_anom = wb.create_sheet("Top Anomalies")
        ws_anom.append(["FABC", "Date", "Producteur", "Commune", "Certification", "Poids livré", "Poids accepté", "Retour", "Taux perte %"])
        for item in data['tracabilite'].get('top_anomalies', []):
            ws_anom.append([
                item.get('numero_fabc'),
                item.get('date_marche'),
                f"{item.get('producteur__code', '')} - {item.get('producteur__nom', '')} {item.get('producteur__prenom', '')}",
                item.get('commune'),
                item.get('certification'),
                item.get('poids_total_livre'),
                item.get('poids_accepte'),
                item.get('poids_retour'),
                item.get('taux_perte_pct'),
            ])
        for cell in ws_anom[1]:
            cell.font = Font(bold=True)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="dashboard_decisionnel_{timestamp}.xlsx"'
        wb.save(response)
        return response

    if export_format == 'pdf':
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas
        except ImportError:
            return Response(
                {'detail': "Le module reportlab n'est pas installé sur le serveur."},
                status=501
            )

        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="dashboard_decisionnel_{timestamp}.pdf"'
        p = canvas.Canvas(response, pagesize=A4)
        y = 800
        p.setFont("Helvetica-Bold", 13)
        p.drawString(40, y, "Dashboard Decisionnel - Resume")
        y -= 30
        p.setFont("Helvetica", 10)

        lines = [
            f"Collecte / Estimation: {data['commercial']['couverture_collecte_pct']}%",
            f"Traçabilité complète: {data['tracabilite']['taux_completude_pct']}%",
            f"Femmes leaders: {data['impact']['taux_femmes_leaders_pct']}%",
            f"Kits couverts: {data['impact']['couverture_kits_pct']}%",
            f"Achats (Ar): {data['commercial']['valeur_achats_ar']}",
            f"Exports (USD): {data['commercial']['valeur_exports_usd']}",
            f"Alertes: {len(data.get('alerts', []))}",
        ]
        for line in lines:
            p.drawString(40, y, line)
            y -= 18

        y -= 10
        p.setFont("Helvetica-Bold", 11)
        p.drawString(40, y, "Top Anomalies")
        y -= 18
        p.setFont("Helvetica", 9)
        for item in data['tracabilite'].get('top_anomalies', [])[:10]:
            text = (
                f"{item.get('numero_fabc')} | {item.get('producteur__code')} | "
                f"Retour={item.get('poids_retour')}kg | Taux={item.get('taux_perte_pct')}%"
            )
            p.drawString(40, y, text[:110])
            y -= 14
            if y < 60:
                p.showPage()
                y = 800
                p.setFont("Helvetica", 9)

        p.showPage()
        p.save()
        return response

    return Response({'detail': 'Format non supporté. Utiliser excel ou pdf.'}, status=400)
